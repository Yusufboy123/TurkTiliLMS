from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document


ROOT = Path(r"D:\Turk tili\Darsliklar")
A1 = ROOT / "A1"
V2 = ROOT / "A1_V2"
OUTPUT = Path(__file__).with_name("a1-v2-assessment.generated.ts")


def clean(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def stable_options(correct: str, candidates: list[str], seed: int) -> list[str]:
    values: list[str] = []
    for value in [correct, *candidates]:
        value = clean(value)
        if value and value.casefold() not in {item.casefold() for item in values}:
            values.append(value)
        if len(values) == 4:
            break
    while len(values) < 3:
        fallback = [
            "Qoida bu shaklni tasdiqlamaydi.",
            "Belgilar boshqa guruhga tegishli.",
            "Bu shakl berilgan vaziyatga mos emas.",
        ][len(values) - 1]
        if fallback.casefold() not in {item.casefold() for item in values}:
            values.append(fallback)
    shift = seed % len(values)
    return values[shift:] + values[:shift]


def structured_distractors(answer: str) -> list[str]:
    swaps = [
        ("qalin", "ingichka"),
        ("Qalin", "Ingichka"),
        ("unli", "undosh"),
        ("Unli", "Undosh"),
        ("DOĞRU", "YANLIŞ"),
    ]
    results: list[str] = []
    for left, right in swaps:
        if left in answer or right in answer:
            marker = "__A1_V2_SWAP__"
            results.append(answer.replace(left, marker).replace(right, left).replace(marker, right))
    chunks = [clean(chunk) for chunk in answer.split(";") if clean(chunk)]
    if len(chunks) > 1:
        results.append("; ".join(chunks[1:] + chunks[:1]))
        results.append("; ".join(reversed(chunks)))
    return results


def parse_l1_practice() -> list[dict[str, object]]:
    document = Document(V2 / "02_MASHQ_DAFTARI" / "01_Alfabe_Interaktiv_Mashqlar_V2.docx")
    items: dict[int, dict[str, str | int]] = {}
    current: dict[str, str | int] | None = None
    section = 0
    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        if not text:
            continue
        if paragraph.style.name.startswith("Heading 1"):
            section += 1
            continue
        match = re.match(r"^(\d+)-mashq\s+([A-Z_]+)$", text)
        if match:
            current = {"number": int(match.group(1)), "source_type": match.group(2), "stage": min(section, 6)}
            items[int(match.group(1))] = current
            continue
        if current is None:
            continue
        for label, key in [
            ("Ko‘rsatma:", "prompt"),
            ("Variantlar yoki bog‘lashlar:", "variants"),
            ("To‘g‘ri javob:", "answer"),
            ("Izoh:", "explanation"),
        ]:
            if text.startswith(label):
                current[key] = clean(text[len(label):])

    result: list[dict[str, object]] = []
    for number in sorted(items):
        raw = items[number]
        prompt = str(raw.get("prompt", ""))
        answer = str(raw.get("answer", ""))
        source_type = str(raw.get("source_type", "CHOICE"))
        variants = [clean(value) for value in str(raw.get("variants", "")).split("|") if clean(value)]
        if not variants:
            variants = re.findall(r"\[([^\]]+)\]", prompt)
        item_type = "TRUE_FALSE" if source_type == "TRUE_FALSE" else "MULTIPLE_CHOICE"
        if item_type == "TRUE_FALSE":
            options = ["To‘g‘ri", "Noto‘g‘ri"]
            answer = "To‘g‘ri" if answer == "DOĞRU" else "Noto‘g‘ri"
        else:
            options = stable_options(answer, variants + structured_distractors(answer), number)
        result.append({
            "id": f"a1v2-d01-p{number:02d}",
            "type": item_type,
            "prompt": f"Ko‘rsatma: {prompt}",
            "options": options,
            "answer": answer,
            "explanation": str(raw.get("explanation", "Darsdagi qoida asosida tekshiring.")),
            "stage": int(raw.get("stage", 1)),
        })
    return result


def student_ready_prompt(day: int, prompt: str) -> str:
    prompt = re.sub(r"^\d+\.\s*", "", clean(prompt))
    special = {
        "Vaziyatga mos iborani yozing: ertalab.": "Vaziyat: Ertalab bir kishini uchratdingiz. Mos turkcha salomlashuvni tanlang.",
        "Vaziyatga mos iborani yozing: umumiy.": "Vaziyat: Kunning istalgan vaqtida bir kishini uchratdingiz. Mos umumiy turkcha salomlashuvni tanlang.",
        "Vaziyatga mos iborani yozing: norasmiy.": "Vaziyat: Yaqin do‘stingiz bilan uchrashdingiz. Mos norasmiy turkcha salomlashuvni tanlang.",
        "Vaziyatga mos iborani yozing: kechqurun.": "Vaziyat: Kechqurun bir kishini uchratdingiz. Mos turkcha salomlashuvni tanlang.",
        "Vaziyatga mos iborani yozing: kelgan mehmon uchun.": "Vaziyat: Mehmon sizning oldingizga keldi. Unga aytiladigan turkcha iborani tanlang.",
    }
    prompt = special.get(prompt, prompt)
    replacements = {
        "yozing": "tanlang",
        "tuzing": "tanlang",
        "ayting": "tanlang",
        "to‘ldiring": "to‘ldiradigan javobni tanlang",
    }
    for source, target in replacements.items():
        prompt = prompt.replace(source, target)
    return f"Ko‘rsatma: To‘g‘ri javobni tanlang.\n{prompt}"


def parse_workbook(day: int, path: Path) -> list[dict[str, object]]:
    document = Document(path)
    answers: dict[int, tuple[str, str]] = {}
    for row in document.tables[-1].rows[1:]:
        cells = [clean(cell.text) for cell in row.cells]
        if cells and cells[0].isdigit():
            answers[int(cells[0])] = (cells[1], cells[2] if len(cells) > 2 else "")

    stage = 0
    questions: list[tuple[int, int, str]] = []
    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        stage_match = re.search(r"([1-7])-BOSQICH", text)
        if stage_match:
            stage = int(stage_match.group(1))
        question_match = re.match(r"^(\d+)\.\s+(.+)$", text)
        if question_match and stage:
            questions.append((int(question_match.group(1)), stage, question_match.group(2)))

    excluded = re.compile(r"MODEL|RUBRIKA|MEZON|erkin|ochiq|45[–-]60|dialog|matn yoz|gap tuz|vaziyat yoz|qayta tekshir", re.I)
    by_stage: dict[int, list[tuple[int, int, str, str, str]]] = defaultdict(list)
    for number, question_stage, prompt in questions:
        answer, answer_type = answers.get(number, ("", ""))
        if question_stage > 5 or not answer or len(answer) > 120 or excluded.search(answer) or excluded.search(prompt):
            continue
        by_stage[question_stage].append((number, question_stage, prompt, answer, answer_type))

    selected: list[tuple[int, int, str, str, str]] = []
    for question_stage in range(1, 6):
        selected.extend(by_stage[question_stage][:5])
    selected.sort(key=lambda item: item[0])

    answer_pool = [item[3] for item in selected]
    result: list[dict[str, object]] = []
    corrections = {(4, 11): "şehre"}
    for index, (number, question_stage, prompt, answer, answer_type) in enumerate(selected, start=1):
        answer = corrections.get((day, number), answer)
        nearby = answer_pool[index:index + 3] + answer_pool[max(0, index - 4):index - 1]
        result.append({
            "id": f"a1v2-d{day:02d}-p{number:02d}",
            "type": "MULTIPLE_CHOICE",
            "prompt": student_ready_prompt(day, prompt),
            "options": stable_options(answer, nearby, day * 100 + number),
            "answer": answer,
            "explanation": f"Darsdagi {answer_type or 'qoida'} bo‘yicha to‘g‘ri javob: {answer}.",
            "stage": question_stage,
        })
    return result


def parse_l1_topic_test() -> list[dict[str, object]]:
    document = Document(V2 / "03_TEST_BANK" / "01_Turk_Alfabesi_ve_Sesler_Topic_Test_V2.docx")
    questions: dict[int, dict[str, str]] = {}
    keys: dict[int, dict[str, str]] = {}
    current: dict[str, str] | None = None
    current_key: dict[str, str] | None = None
    for paragraph in document.paragraphs:
        text = clean(paragraph.text)
        question_match = re.match(r"^(\d+)-savol\s+([A-Z_]+)$", text)
        key_match = re.match(r"^(\d+)-savol kaliti$", text)
        if question_match:
            number = int(question_match.group(1))
            current = {"type": question_match.group(2)}
            questions[number] = current
            current_key = None
            continue
        if key_match:
            number = int(key_match.group(1))
            current_key = {}
            keys[number] = current_key
            current = None
            continue
        if current is not None and text:
            if text.startswith("Variantlar yoki topshiriq:"):
                current["variants"] = clean(text.split(":", 1)[1])
            elif "prompt" not in current:
                current["prompt"] = text
        if current_key is not None:
            if text.startswith("To‘g‘ri javob:"):
                current_key["answer"] = clean(text.split(":", 1)[1])
            elif text.startswith("Izoh:"):
                current_key["explanation"] = clean(text.split(":", 1)[1])

    result: list[dict[str, object]] = []
    for number in sorted(questions):
        question = questions[number]
        key = keys[number]
        source_type = question["type"]
        variants = [clean(value) for value in question.get("variants", "").split("|") if clean(value)]
        if not variants:
            variants = re.findall(r"\[([^\]]+)\]", question.get("prompt", ""))
        answer = key["answer"]
        question_type = "TRUE_FALSE" if source_type == "TRUE_FALSE" else "MULTIPLE_CHOICE"
        if question_type == "TRUE_FALSE":
            options = ["To‘g‘ri", "Noto‘g‘ri"]
            answer = "To‘g‘ri" if answer == "DOĞRU" else "Noto‘g‘ri"
        else:
            options = stable_options(answer, variants + structured_distractors(answer), 1000 + number)
        result.append({
            "type": question_type,
            "prompt": question["prompt"],
            "explanation": key["explanation"],
            "points": 1,
            "position": number,
            "options": [
                {"text": option, "isCorrect": option == answer, "position": position}
                for position, option in enumerate(options, start=1)
            ],
        })
    return result


def parse_test_bank() -> dict[int, list[dict[str, object]]]:
    document = Document(A1 / "03_TEST_BANK" / "A1_TEST_BANK.docx")
    prompts: dict[tuple[int, int], str] = {}
    for paragraph in document.paragraphs:
        match = re.match(r"^(\d+)\.(\d+) \[MULTIPLE CHOICE\] (.+)$", clean(paragraph.text))
        if match:
            prompts[(int(match.group(1)), int(match.group(2)))] = match.group(3)

    option_tables = document.tables[2:218]
    options: dict[tuple[int, int], list[str]] = {}
    for offset, table in enumerate(option_tables):
        day = offset // 18 + 1
        number = offset % 18 + 1
        options[(day, number)] = [clean(cell.text) for cell in table.rows[1].cells]

    keys: dict[tuple[int, int], tuple[str, str]] = {}
    for day, table in enumerate(document.tables[218:230], start=1):
        for row in table.rows[1:]:
            cells = [clean(cell.text) for cell in row.cells]
            number = int(cells[0].split(".")[1])
            correct = cells[1].split("—", 1)[1].strip() if "—" in cells[1] else cells[1]
            keys[(day, number)] = (correct, cells[2])

    selected_numbers = [1, 2, 3, 7, 8, 9, 13, 14, 15, 16, 17, 18]
    result: dict[int, list[dict[str, object]]] = {}
    for day in range(2, 13):
        rows: list[dict[str, object]] = []
        for position, number in enumerate(selected_numbers, start=1):
            correct, explanation = keys[(day, number)]
            question_options = options[(day, number)]
            if correct not in question_options:
                raise ValueError(f"Day {day} question {number}: key is not in options")
            rows.append({
                "type": "MULTIPLE_CHOICE",
                "prompt": prompts[(day, number)],
                "explanation": explanation,
                "points": 1,
                "position": position,
                "options": [
                    {"text": option, "isCorrect": option == correct, "position": option_position}
                    for option_position, option in enumerate(question_options, start=1)
                ],
            })
        result[day] = rows
    return result


def ts(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, indent=2).replace("\\u2028", "\\u2028").replace("\\u2029", "\\u2029")


def main() -> None:
    workbooks = sorted(path for path in (A1 / "02_MASHQ_DAFTARI").glob("*.docx") if not path.name.startswith("~$"))
    practice: dict[int, list[dict[str, object]]] = {1: parse_l1_practice()}
    for day, path in enumerate(workbooks[1:], start=2):
        practice[day] = parse_workbook(day, path)
    topic = parse_test_bank()
    topic[1] = parse_l1_topic_test()
    ordered_topic = {day: topic[day] for day in range(1, 13)}
    output = """// Generated from the approved A1 V2 lesson-one package and authoritative A1 workbooks/test bank.\n// Regenerate with generate-a1-v2-assessment.py; do not hand-edit question keys.\nimport type { A1PracticeItemDefinition, A1QuestionDefinition } from './a1-content.js';\n\n"""
    output += f"export const a1V2PracticeByDay: Readonly<Record<number, A1PracticeItemDefinition[]>> = {ts(practice)};\n\n"
    output += f"export const a1V2QuestionsByDay: Readonly<Record<number, A1QuestionDefinition[]>> = {ts(ordered_topic)};\n"
    OUTPUT.write_text(output, encoding="utf-8", newline="\n")
    print("practice", {day: len(items) for day, items in practice.items()})
    print("topic", {day: len(items) for day, items in ordered_topic.items()})


if __name__ == "__main__":
    main()
