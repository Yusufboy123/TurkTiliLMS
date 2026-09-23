"""Generate A2 V2 theory, deterministic practice, and topic-test sources.

The authoritative inputs are the A2 DOCX package. Vocabulary is deliberately
excluded: teacher-provided vocabulary rows are owned by the existing vocabulary
pipeline and must never be reconciled by this generator.
"""

from __future__ import annotations

import json
import os
import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


HERE = Path(__file__).resolve().parent
SOURCE_ROOT = Path(os.environ.get("A2_SOURCE_ROOT", r"D:\Turk tili\Darsliklar\A2"))
THEORY_DIR = SOURCE_ROOT / "01_DARSLIK"
PRACTICE_DIR = SOURCE_ROOT / "02_MASHQ_DAFTARI"
TEST_BANK = SOURCE_ROOT / "03_TEST_BANK" / "A2_TEST_BANK.docx"


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def iter_blocks(document: DocumentType) -> Iterable[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def markdown_table(table: Table) -> str:
    rows = [
        [normalize(cell.text).replace("|", "\\|") for cell in row.cells]
        for row in table.rows
    ]
    if not rows or not any(any(cell for cell in row) for row in rows):
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    header = rows[0]
    body = rows[1:]
    if not body:
        return "| " + " | ".join(header) + " |\n| " + " | ".join(["---"] * width) + " |"
    return "\n".join(
        [
            "| " + " | ".join(header) + " |",
            "| " + " | ".join(["---"] * width) + " |",
            *("| " + " | ".join(row) + " |" for row in body),
        ]
    )


def extract_theory(path: Path) -> str:
    document = Document(path)
    output: list[str] = []
    skipping_vocabulary = False
    heading_seen = False

    for block in iter_blocks(document):
        if isinstance(block, Paragraph):
            text = normalize(block.text)
            if not text:
                continue
            style = (block.style.name or "").lower()
            is_heading = style.startswith("heading")

            if text == "Dars lug‘ati: NEW va REVIEW":
                skipping_vocabulary = True
                continue
            if skipping_vocabulary:
                if is_heading and text == "Dars xulosasi":
                    skipping_vocabulary = False
                else:
                    continue

            if text in {
                "A2 DARSLIK",
                "Mustaqil o‘rganish • dars • mashq • takrorlash",
            } or re.fullmatch(r"\d+-DARS", text):
                continue

            if is_heading:
                prefix = "##" if not heading_seen else "###"
                heading_seen = True
                output.append(f"{prefix} {text}")
            elif "list" in style:
                output.append(f"- {text}")
            else:
                output.append(text)
        else:
            if skipping_vocabulary:
                continue
            cells = [normalize(cell.text) for row in block.rows for cell in row.cells]
            if not cells:
                continue
            if cells[0] == "Daraja" or "Source ID" in cells:
                continue
            rendered = markdown_table(block)
            if rendered:
                output.append(rendered)

    theory = "\n\n".join(output).strip()
    forbidden = ("Source ID", "Dars lug‘ati: NEW va REVIEW", "A2 DARSLIK")
    if len(theory) < 2_000 or any(marker in theory for marker in forbidden):
        raise ValueError(f"Theory extraction failed quality checks: {path.name}")
    return theory


def workbook_items(path: Path) -> tuple[dict[int, str], dict[int, str]]:
    document = Document(path)
    answers: dict[int, str] = {}
    for row in document.tables[-1].rows[1:]:
        number = normalize(row.cells[0].text)
        if number.isdigit():
            answers[int(number)] = normalize(row.cells[1].text)

    prompts: dict[int, str] = {}
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        heading = normalize(paragraph.text)
        match = re.match(r"^(\d+)-mashq\.", heading)
        if not match:
            continue
        number = int(match.group(1))
        cursor = index + 1
        candidate = ""
        while cursor < len(paragraphs):
            current = normalize(paragraphs[cursor].text)
            style = (paragraphs[cursor].style.name or "").lower()
            if current and style.startswith("heading"):
                break
            numbered = re.match(rf"^{number}\.\s*(.+)$", current)
            if numbered:
                candidate = numbered.group(1).strip()
                break
            cursor += 1
        if candidate:
            prompts[number] = candidate
    return prompts, answers


def unique_source_numbers(answers: dict[int, str], start: int, end: int, transform=lambda value: value) -> list[int]:
    result: list[int] = []
    seen: set[str] = set()
    for number in range(start, end + 1):
        value = transform(answers[number])
        key = normalize(value).casefold()
        if key not in seen:
            seen.add(key)
            result.append(number)
    return result


def translated_answer(value: str) -> str:
    marker = "manba tarjima:"
    if marker not in value:
        raise ValueError(f"Translation marker missing: {value}")
    return value.split(marker, 1)[1].strip()


def corrected_answer(value: str) -> str:
    return re.split(r"\s+[—–]\s+", value, maxsplit=1)[0].strip()


def option_set(correct: str, pool: list[str], seed: int) -> list[str]:
    correct_key = normalize(correct).casefold()
    distractors: list[str] = []
    seen: set[str] = set()
    for value in pool:
        key = normalize(value).casefold()
        if key == correct_key or key in seen:
            continue
        seen.add(key)
        distractors.append(value)
    if len(distractors) < 3:
        raise ValueError(f"Not enough distinct distractors for {correct!r}")
    start = seed % len(distractors)
    chosen = [distractors[(start + offset) % len(distractors)] for offset in range(3)]
    options = chosen[:]
    options.insert(seed % 4, correct)
    if len({normalize(value).casefold() for value in options}) != 4:
        raise ValueError(f"Duplicate options generated for {correct!r}")
    return options


def make_practice(day: int, path: Path) -> list[dict[str, Any]]:
    prompts, answers = workbook_items(path)
    required = set(range(1, 41))
    if not required.issubset(prompts) or not required.issubset(answers):
        raise ValueError(f"Workbook 1-40 incomplete: {path.name}")

    items: list[dict[str, Any]] = []

    def append(number: int, stage: int, prompt: str, answer: str, explanation: str, pool: list[str]) -> None:
        items.append(
            {
                "id": f"a2v2-d{day:02d}-p{number:03d}",
                "type": "MULTIPLE_CHOICE",
                "prompt": prompt,
                "options": option_set(answer, pool, day * 41 + number),
                "answer": answer,
                "explanation": explanation,
                "stage": stage,
            }
        )

    translation_numbers = list(range(1, 9))
    translations = [translated_answer(answers[number]) for number in translation_numbers]
    for number in translation_numbers:
        answer = translated_answer(answers[number])
        append(
            number,
            1,
            f"Ko‘rsatma: Turkcha gapning o‘zbekcha ma’nosini tanlang.\nTurkcha gap: {prompts[number]}",
            answer,
            f"To‘g‘ri ma’no: {answer}",
            translations,
        )

    table_numbers = unique_source_numbers(answers, 9, 16)
    table_answers = [answers[number] for number in table_numbers]
    for number in table_numbers:
        anchor = answers[number].split(" / ", 1)[0].strip()
        append(
            number,
            2,
            f"Ko‘rsatma: Dars jadvalidagi to‘liq va to‘g‘ri qatorni tanlang.\nTayanch: {anchor}",
            answers[number],
            f"Dars jadvalidagi mos qator: {answers[number]}",
            table_answers,
        )

    missing_numbers = list(range(17, 25))
    missing_answers = [answers[number] for number in missing_numbers]
    for number in missing_numbers:
        append(
            number,
            3,
            f"Ko‘rsatma: Bo‘shliqni to‘ldiradigan mos birlikni tanlang.\n{prompts[number]}",
            answers[number],
            f"Gap mazmuniga mos birlik: {answers[number]}",
            missing_answers,
        )

    transformation_numbers = unique_source_numbers(answers, 25, 32)
    transformation_answers = [answers[number] for number in transformation_numbers]
    for number in transformation_numbers:
        append(
            number,
            4,
            f"Ko‘rsatma: Berilgan vazifaga mos turkcha shaklni tanlang.\n{prompts[number]}",
            answers[number],
            f"Dars qolipiga mos shakl: {answers[number]}",
            transformation_answers,
        )

    correction_numbers = unique_source_numbers(answers, 33, 40, corrected_answer)
    correction_answers = [corrected_answer(answers[number]) for number in correction_numbers]
    for number in correction_numbers:
        answer = corrected_answer(answers[number])
        append(
            number,
            5,
            f"Ko‘rsatma: Noto‘g‘ri gapning tuzatilgan variantini tanlang.\nNoto‘g‘ri gap: {prompts[number]}",
            answer,
            answers[number],
            correction_answers,
        )

    if len(items) < 30:
        raise ValueError(f"Too few deterministic practice items for day {day}: {len(items)}")
    return items


def parse_test_bank(path: Path) -> dict[int, list[dict[str, Any]]]:
    document = Document(path)
    keys: dict[tuple[int, int], dict[str, str]] = {}
    for table in document.tables:
        if not table.rows or normalize(table.rows[0].cells[0].text) != "№":
            continue
        for row in table.rows[1:]:
            label = normalize(row.cells[0].text)
            match = re.fullmatch(r"(\d+)\.(\d+)", label)
            if not match:
                continue
            day, position = map(int, match.groups())
            keys[(day, position)] = {
                "answer": normalize(row.cells[1].text),
                "explanation": normalize(row.cells[2].text),
            }

    parsed: dict[int, list[dict[str, Any]]] = {day: [] for day in range(1, 15)}
    sequence = list(iter_blocks(document))
    for index, block in enumerate(sequence):
        if not isinstance(block, Paragraph):
            continue
        text = normalize(block.text)
        match = re.match(r"^(\d+)\.(\d+) \[([^\]]+)\]\s*(.+)$", text)
        if not match:
            continue
        day, position = int(match.group(1)), int(match.group(2))
        if position > 25:
            continue
        kind, prompt = match.group(3), match.group(4)
        key = keys.get((day, position))
        if not key:
            raise ValueError(f"Test key missing for {day}.{position}")
        record: dict[str, Any] = {
            "sourcePosition": position,
            "kind": kind,
            "prompt": prompt,
            "key": key,
        }
        if kind == "MULTIPLE CHOICE":
            next_table = next((candidate for candidate in sequence[index + 1 : index + 4] if isinstance(candidate, Table)), None)
            if next_table is None:
                raise ValueError(f"Options missing for {day}.{position}")
            record["options"] = [normalize(cell.text) for cell in next_table.rows[1].cells]
        parsed[day].append(record)

    for day, questions in parsed.items():
        if len(questions) != 25:
            raise ValueError(f"Expected 25 objective test questions for day {day}, got {len(questions)}")
    return parsed


def make_questions(day: int, records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    missing_pool = [record["key"]["answer"] for record in records if record["kind"] == "MISSING WORD"]
    questions: list[dict[str, Any]] = []
    letters = {"A": 0, "B": 1, "C": 2, "D": 3}

    for target_position, record in enumerate(records, start=1):
        kind = record["kind"]
        raw_answer = record["key"]["answer"]
        if kind == "MULTIPLE CHOICE":
            letter = raw_answer.split(" ", 1)[0]
            if letter not in letters:
                raise ValueError(f"Invalid MCQ key for day {day}: {raw_answer}")
            options = record["options"]
            correct_index = letters[letter]
            if len(options) != 4:
                raise ValueError(f"Invalid MCQ options for day {day}.{target_position}")
            keyed_text = re.split(r"\s+[—–-]\s+", raw_answer, maxsplit=1)
            if len(keyed_text) != 2 or normalize(keyed_text[1]).casefold() != normalize(options[correct_index]).casefold():
                raise ValueError(f"MCQ key text does not match its option for day {day}.{target_position}")
        elif kind == "TRUE/FALSE":
            options = ["To‘g‘ri", "Noto‘g‘ri"]
            correct_index = 0 if raw_answer.upper().startswith(("DOĞRU", "TO‘G‘RI")) else 1
        elif kind == "MISSING WORD":
            correct = raw_answer
            options = option_set(correct, missing_pool, day * 37 + target_position)
            correct_index = options.index(correct)
        else:
            raise ValueError(f"Unsupported objective question type: {kind}")

        questions.append(
            {
                "type": "TRUE_FALSE" if kind == "TRUE/FALSE" else "MULTIPLE_CHOICE",
                "prompt": record["prompt"],
                "explanation": f"To‘g‘ri javob: {options[correct_index]}",
                "points": 1,
                "position": target_position,
                "options": [
                    {"text": option, "isCorrect": index == correct_index, "position": index + 1}
                    for index, option in enumerate(options)
                ],
            }
        )
    return questions


def assert_quality(
    theory: dict[int, str],
    practice: dict[int, list[dict[str, Any]]],
    questions: dict[int, list[dict[str, Any]]],
) -> None:
    forbidden = re.compile(
        r"\bMODEL\s*:|\bRUBRIKA\s*:|\bPLACEHOLDER\b|Ochiq javob|Muqobil tabiiy javoblar|Mezon:",
        re.I,
    )
    for day in range(1, 15):
        if day not in theory or day not in practice or day not in questions:
            raise ValueError(f"Day {day} is incomplete")
        practice_prompts: set[str] = set()
        for item in practice[day]:
            if forbidden.search(item["prompt"]) or forbidden.search(item["answer"]):
                raise ValueError(f"Open/placeholder practice item: {item['id']}")
            options = item.get("options", [])
            if len(options) != len(set(options)) or options.count(item["answer"]) != 1:
                raise ValueError(f"Invalid practice options: {item['id']}")
            normalized_prompt = normalize(item["prompt"]).casefold()
            if normalized_prompt in practice_prompts:
                raise ValueError(f"Duplicate practice prompt: {item['id']}")
            practice_prompts.add(normalized_prompt)

        topic_prompts: set[str] = set()
        correct_positions: set[int] = set()
        for question in questions[day]:
            if forbidden.search(question["prompt"]):
                raise ValueError(f"Open/placeholder topic question: day {day}")
            option_texts = [option["text"] for option in question["options"]]
            correct = [index for index, option in enumerate(question["options"]) if option["isCorrect"]]
            if len(option_texts) != len(set(option_texts)) or len(correct) != 1:
                raise ValueError(f"Invalid topic options: day {day}, position {question['position']}")
            correct_positions.add(correct[0])
            normalized_prompt = normalize(question["prompt"]).casefold()
            question_identity = "\u241f".join(
                [normalized_prompt, *sorted(normalize(value).casefold() for value in option_texts)]
            )
            if question_identity in topic_prompts:
                raise ValueError(f"Duplicate topic prompt: day {day}, position {question['position']}")
            topic_prompts.add(question_identity)
        if len(correct_positions) < 2:
            raise ValueError(f"Correct answer position bias detected for day {day}")


def write_typescript(path: Path, header: str, export_name: str, payload: dict[int, Any], type_name: str) -> None:
    serialized = json.dumps(payload, ensure_ascii=False, indent=2)
    path.write_text(
        f"// {header}\n"
        "// Regenerate with generate-a2-v2-content.py; do not hand-edit.\n"
        f"import type {{ {type_name} }} from './a2-content.js';\n\n"
        f"export const {export_name}: Readonly<Record<number, {type_name}[]>> = {serialized};\n",
        encoding="utf-8",
        newline="\n",
    )


def main() -> None:
    theory_files = sorted(THEORY_DIR.glob("*.docx"))
    practice_files = sorted(PRACTICE_DIR.glob("*.docx"))
    if len(theory_files) != 14 or len(practice_files) != 14 or not TEST_BANK.exists():
        raise FileNotFoundError("The authoritative A2 package is incomplete")

    theory = {day: extract_theory(path) for day, path in enumerate(theory_files, start=1)}
    practice = {day: make_practice(day, path) for day, path in enumerate(practice_files, start=1)}
    parsed_tests = parse_test_bank(TEST_BANK)
    questions = {day: make_questions(day, parsed_tests[day]) for day in range(1, 15)}
    assert_quality(theory, practice, questions)

    theory_payload = json.dumps(theory, ensure_ascii=False, indent=2)
    (HERE / "a2-v2-theory.generated.ts").write_text(
        "// Generated from the authoritative A2 theory DOCX files.\n"
        "// Regenerate with generate-a2-v2-content.py; do not hand-edit.\n"
        f"export const a2V2TheoryByDay: Readonly<Record<number, string>> = {theory_payload};\n",
        encoding="utf-8",
        newline="\n",
    )
    write_typescript(
        HERE / "a2-v2-practice.generated.ts",
        "Generated from the authoritative A2 workbook DOCX files; open responses are intentionally excluded from exact grading.",
        "a2V2PracticeByDay",
        practice,
        "A2PracticeItemDefinition",
    )
    write_typescript(
        HERE / "a2-v2-questions.generated.ts",
        "Generated from the authoritative A2 test bank; only objectively gradable questions 1-25 are included.",
        "a2V2QuestionsByDay",
        questions,
        "A2QuestionDefinition",
    )
    print(
        json.dumps(
            {
                "lessons": 14,
                "practiceByDay": {day: len(items) for day, items in practice.items()},
                "practiceTotal": sum(map(len, practice.values())),
                "questionsByDay": {day: len(items) for day, items in questions.items()},
                "questionsTotal": sum(map(len, questions.values())),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
